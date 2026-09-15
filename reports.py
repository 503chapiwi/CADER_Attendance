"""Generación del Listado de Beneficiarios y el Informe Mensual a partir de los datos de Kobo."""

import io
from datetime import date, datetime
from pathlib import Path

import openpyxl

from kobo_client import _slug, label_for

PLANTILLAS = Path(__file__).parent / "plantillas"

MESES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

# categoria (Kobo) -> hoja del listado de beneficiarios
CATEGORIA_HOJA = {
    "agricola": "AGRÍCOLA",
    "pecuario": "PECUARIO",
    "hogar_rural": "HOGAR RURAL",
}

# Fila del título de cada sección municipal en la plantilla del informe;
# los eventos van en las 10 filas siguientes.
SECCION_INFORME = {
    "momostenango": 11,
    "san_andres_xecul": 24,
    "san_bartolo": 37,
    "san_cristobal_totonicapan": 50,
    "san_francisco_el_alto": 63,
    "santa_lucia_la_reforma": 76,
    "santa_maria_chiquimula": 89,
    "totonicapan": 102,
}

# orientacion (Kobo) -> columna de la X en el informe (I..P = 9..16)
ORIENTACION_COL = {
    "agricola": 9,
    "pecuario": 10,
    "forestal": 11,
    "hidrobiologico": 12,
    "hogar_rural": 13,
    "juventud_y_ninez": 14,
    "asociatividad_y_comercializacion": 15,
    "otro": 16,
}

SEXO_MAP = {"hombre": "Masculino", "masculino": "Masculino", "mujer": "Femenino", "feminino": "Femenino", "femenino": "Femenino"}
PUEBLO_MAP = {"maya": "Maya", "ladino": "Ladino", "xinca": "Xinca", "garifuna": "Garifuna", "afrodescendiente": "Afrodescendiente", "extranjero": "Extranjero"}
COMUNIDAD_LING_MAP = {
    "kiche": "K´iche´", "k_iche": "K´iche´", "mam": "Mam", "kakchiquel": "Kaqchikel",
    "kaqchikel": "Kaqchikel", "ixil": "Ixil", "tzutujil": "Tz´utujil", "tz_utujil": "Tz´utujil",
    "qeqchi": "Q´eqchi´", "q_eqchi": "Q´eqchi´", "uspanteca": "Uspanteka", "uspanteka": "Uspanteka",
    "espanol": "Idioma Español", "idioma_espanol": "Idioma Español",
}
DISCAPACIDAD_MAP = {"ninguna": "No tiene discapacidad", "visual": "Visual", "auditiva": "Auditiva", "fisica": "Física", "motora": "Física", "intelectual": "Intelectual"}
SI_NO_MAP = {"si": "Si", "no": "No"}


def _clean(value):
    if value is None or value != value:  # None o NaN de pandas
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("nan", "none") else text


def _to_int(value):
    try:
        return int(float(str(value)))
    except (TypeError, ValueError):
        return None


def _parse_fecha(sub):
    try:
        return datetime.strptime(_clean(sub.get("fecha"))[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def filtrar_envios(submissions, municipio_code, mes, anio):
    """Envíos del municipio y mes/año elegidos (mes 1-12).
    Con municipio_code=None se incluyen todos los municipios."""
    muni_slug = _slug(municipio_code) if municipio_code else None
    out = []
    for sub in submissions:
        f = _parse_fecha(sub)
        if f is None or f.month != mes or f.year != anio:
            continue
        # _slug también cubre datos de versiones viejas que traían la etiqueta
        # ("San Cristóbal Totonicapán") en vez del código
        if muni_slug and _slug(sub.get("municipio")) != muni_slug:
            continue
        out.append(sub)
    return out


def registro_nuevos_miembros(submissions):
    """Personas registradas como nuevas en cualquier envío del formulario
    (repeats 'nuevos' y 'nuevo_grupo'), indexadas por CUI. Sirven de respaldo
    cuando un CUI no aparece en la base de participantes."""
    registro = {}
    for sub in submissions:
        for repeat_key, prefix in (("nuevos", "n_"), ("nuevo_grupo", "ng_")):
            for item in sub.get(repeat_key) or []:
                cui = _clean(item.get(prefix + "cui"))
                if not cui:
                    continue
                registro[cui] = {
                    "primer_apellido": _clean(item.get(prefix + "primer_apellido")),
                    "segundo_apellido": _clean(item.get(prefix + "segundo_apellido")),
                    "primer_nombre": _clean(item.get(prefix + "primer_nombre")),
                    "segundo_nombre": _clean(item.get(prefix + "segundo_nombre")),
                    "sexo": SEXO_MAP.get(_slug(item.get(prefix + "sexo")), ""),
                    "anio": _to_int(item.get(prefix + "anio")),
                    "pueblo": PUEBLO_MAP.get(_slug(item.get(prefix + "pueblo")), ""),
                    "comunidad_ling": COMUNIDAD_LING_MAP.get(_slug(item.get(prefix + "comunidad_ling")), ""),
                    "discapacidad": DISCAPACIDAD_MAP.get(_slug(item.get(prefix + "discapacidad")), ""),
                    "retornado": SI_NO_MAP.get(_slug(item.get(prefix + "retornado")), ""),
                }
    return registro


def indexar_participantes(df):
    """Base de participantes CADER (CSV de Kobo o archivo subido) -> dict por CUI."""
    idx = {}
    cui_col = "name" if "name" in df.columns else next((c for c in df.columns if _slug(c) in ("cui", "name", "codigo_unico_de_identificacion")), None)
    if cui_col is None:
        raise ValueError("La base de participantes no tiene columna de CUI ('name' o 'cui').")
    for _, row in df.iterrows():
        cui = _clean(row.get(cui_col))
        if not cui:
            continue
        discapacidades = [
            nombre for col, nombre in (
                ("discapacidad_visual", "Visual"),
                ("discapacidad_auditiva", "Auditiva"),
                ("discapacidad_motora", "Física"),
                ("discapacidad_intelectual", "Intelectual"),
            ) if _clean(row.get(col)).upper() in ("X", "SI", "SÍ", "1", "TRUE")
        ]
        if len(discapacidades) == 0:
            discapacidad = "No tiene discapacidad"
        elif len(discapacidades) == 1:
            discapacidad = discapacidades[0]
        else:
            discapacidad = "Múltiple"
        idx[cui] = {
            "primer_apellido": _clean(row.get("primer_apellido")),
            "segundo_apellido": _clean(row.get("segundo_apellido")),
            "primer_nombre": _clean(row.get("primer_nombre")),
            "segundo_nombre": _clean(row.get("segundo_nombre")),
            "sexo": SEXO_MAP.get(_slug(row.get("genero")), _clean(row.get("genero"))),
            "anio": _to_int(row.get("ano_nac")),
            "pueblo": PUEBLO_MAP.get(_slug(row.get("pueblo")), _clean(row.get("pueblo")).title()),
            "comunidad_ling": COMUNIDAD_LING_MAP.get(_slug(row.get("comunidad_linguistica")), _clean(row.get("comunidad_linguistica"))),
            "discapacidad": discapacidad,
            "retornado": SI_NO_MAP.get(_slug(row.get("retornado")), _clean(row.get("retornado")).title()),
        }
    return idx


def indexar_promotores(df):
    """cader_id -> nombre completo del promotor."""
    idx = {}
    for _, row in df.iterrows():
        cader = _clean(row.get("name"))
        if not cader:
            continue
        partes = [_clean(row.get(c)) for c in ("primer_nombre", "segundo_nombre", "primer_apellido", "segundo_apellido")]
        idx[_slug(cader)] = " ".join(p for p in partes if p)
    return idx


def generar_listado(envios, participantes_idx, nuevos_idx, choice_maps, field_lists, mes, anio):
    """Llena la plantilla del listado de beneficiarios. Devuelve (bytes, resumen)."""
    # beneficiarios[hoja][cui] = {"capacitaciones": [...], "municipio": ..., "comunidad": ...}
    beneficiarios = {hoja: {} for hoja in CATEGORIA_HOJA.values()}
    eventos = 0
    for sub in envios:
        ldb = _slug(sub.get("listado_de_beneficiarios"))
        if ldb == "no":
            continue
        # versiones viejas del formulario no tenían esta pregunta: se incluyen
        # si traen lista de asistentes
        if ldb != "si" and not _clean(sub.get("asistentes")):
            continue
        hoja = CATEGORIA_HOJA.get(_slug(sub.get("categoria")))
        if hoja is None:
            continue
        eventos += 1
        capacitacion = label_for(choice_maps, field_lists, "capacitacion", sub.get("capacitacion"))
        municipio = label_for(choice_maps, field_lists, "municipio", sub.get("municipio"))
        comunidad = label_for(choice_maps, field_lists, "comunidad", sub.get("comunidad"))
        cuis = _clean(sub.get("asistentes")).split()
        # los miembros nuevos registrados en este envío también asistieron
        for repeat_key, prefix in (("nuevos", "n_"), ("nuevo_grupo", "ng_")):
            cuis += [_clean(i.get(prefix + "cui")) for i in (sub.get(repeat_key) or []) if _clean(i.get(prefix + "cui"))]
        for cui in cuis:
            b = beneficiarios[hoja].setdefault(cui, {"capacitaciones": [], "municipio": municipio, "comunidad": comunidad})
            if capacitacion and capacitacion not in b["capacitaciones"]:
                b["capacitaciones"].append(capacitacion)

    wb = openpyxl.load_workbook(PLANTILLAS / "plantilla_listado.xlsx")
    sin_datos = []
    total = 0
    for hoja, personas in beneficiarios.items():
        ws = wb[hoja]
        ws["C4"] = "TOTONICAPÁN"
        ws["P3"] = MESES[mes - 1]
        ws["P5"] = anio
        ws["O1"] = f"Fecha de reporte: {date.today().strftime('%d/%m/%Y')}"
        fila = 8
        for cui, b in sorted(personas.items(), key=lambda kv: (kv[1]["municipio"], kv[0])):
            datos = participantes_idx.get(cui) or nuevos_idx.get(cui)
            if datos is None:
                datos = {}
                sin_datos.append({"CUI": cui, "Hoja": hoja, "Comunidad": b["comunidad"]})
            ws.cell(fila, 1).value = cui
            ws.cell(fila, 2).value = datos.get("primer_apellido", "")
            ws.cell(fila, 3).value = datos.get("segundo_apellido", "")
            ws.cell(fila, 4).value = datos.get("primer_nombre", "")
            ws.cell(fila, 5).value = datos.get("segundo_nombre", "")
            ws.cell(fila, 6).value = datos.get("sexo", "")
            ws.cell(fila, 7).value = datos.get("anio")
            ws.cell(fila, 8).value = "Totonicapán"
            ws.cell(fila, 9).value = b["municipio"]
            ws.cell(fila, 10).value = b["comunidad"]
            ws.cell(fila, 11).value = datos.get("pueblo", "")
            ws.cell(fila, 12).value = datos.get("comunidad_ling", "")
            ws.cell(fila, 13).value = datos.get("discapacidad", "")
            ws.cell(fila, 14).value = datos.get("retornado", "")
            ws.cell(fila, 15).value = "; ".join(b["capacitaciones"])
            fila += 1
            total += 1

    buf = io.BytesIO()
    wb.save(buf)
    resumen = {
        "eventos": eventos,
        "beneficiarios": total,
        "por_hoja": {h: len(p) for h, p in beneficiarios.items()},
        "sin_datos": sin_datos,
    }
    return buf.getvalue(), resumen


def _llenar_evento(ws, fila, n, sub, promotores_idx, choice_maps, field_lists):
    comunidad = label_for(choice_maps, field_lists, "comunidad", sub.get("comunidad"))
    municipio = label_for(choice_maps, field_lists, "municipio", sub.get("municipio"))
    fecha_texto = _clean(sub.get("fecha_diferente"))
    if not fecha_texto:
        f = _parse_fecha(sub)
        fecha_texto = f.strftime("%d/%m/%Y") if f else ""
    ws.cell(fila, 1).value = n
    ws.cell(fila, 2).value = _clean(sub.get("nombre_evento"))
    ws.cell(fila, 3).value = fecha_texto
    ws.cell(fila, 4).value = ", ".join(p for p in (comunidad, municipio) if p)
    ws.cell(fila, 5).value = _to_int(sub.get("duracion"))
    ws.cell(fila, 6).value = promotores_idx.get(_slug(sub.get("cader_id")), "")
    ws.cell(fila, 7).value = _to_int(sub.get("telefono_promotor")) or _clean(sub.get("telefono_promotor"))
    ws.cell(fila, 8).value = _clean(sub.get("nombre_capacitador"))
    col_x = ORIENTACION_COL.get(_slug(sub.get("orientacion")))
    if col_x:
        ws.cell(fila, col_x).value = "X"
    ws.cell(fila, 17).value = _to_int(sub.get("h_promotores")) or 0
    ws.cell(fila, 18).value = _to_int(sub.get("m_promotores")) or 0
    ws.cell(fila, 19).value = _to_int(sub.get("otros_part_h")) or 0
    ws.cell(fila, 20).value = _to_int(sub.get("otros_part_m")) or 0
    ws.cell(fila, 23).value = _clean(sub.get("institucion"))
    ws.cell(fila, 24).value = _to_int(sub.get("costo"))
    verificacion = {_slug(v) for v in _clean(sub.get("verificacion")).split()}
    ws.cell(fila, 25).value = "X" if "fotos" in verificacion else None
    ws.cell(fila, 26).value = "X" if "listados" in verificacion else None
    ws.cell(fila, 27).value = "X" if "otros" in verificacion else None
    ws.cell(fila, 28).value = _clean(sub.get("observaciones"))


def generar_informe(envios, promotores_idx, choice_maps, field_lists, municipio_code, mes, anio):
    """Llena la plantilla del informe mensual. Con un municipio llena solo su
    sección; con municipio_code=None llena las secciones de todos los
    municipios que tengan eventos. Devuelve (bytes, resumen)."""
    if municipio_code is None:
        munis = list(SECCION_INFORME)
        encabezado = "TODO EL DEPARTAMENTO"
    else:
        muni_slug = _slug(municipio_code)
        if muni_slug not in SECCION_INFORME:
            raise ValueError(f"Municipio no reconocido: {municipio_code}")
        munis = [muni_slug]
        encabezado = label_for(choice_maps, field_lists, "municipio", municipio_code).upper()

    wb = openpyxl.load_workbook(PLANTILLAS / "plantilla_informe.xlsx")
    ws = wb["FORMATO DE CAPACITACIONES 2026"]
    ws.cell(6, 6).value = MESES[mes - 1].upper()
    ws.cell(6, 17).value = encabezado

    total = 0
    advertencias = []
    for muni in munis:
        eventos = [
            s for s in envios
            if _slug(s.get("informe")) == "si" and _slug(s.get("municipio")) == muni
        ]
        eventos.sort(key=lambda s: (_parse_fecha(s) or date.min))
        if len(eventos) > 10:
            etiqueta = label_for(choice_maps, field_lists, "municipio", muni)
            advertencias.append(
                f"{etiqueta} tiene {len(eventos)} eventos con informe este mes, pero la "
                f"plantilla solo tiene 10 filas por municipio. Se incluyeron los primeros 10 (por fecha)."
            )
        for n, sub in enumerate(eventos[:10], start=1):
            _llenar_evento(ws, SECCION_INFORME[muni] + n, n, sub, promotores_idx, choice_maps, field_lists)
        total += min(len(eventos), 10)

    buf = io.BytesIO()
    wb.save(buf)
    resumen = {"eventos": total, "advertencias": advertencias}
    return buf.getvalue(), resumen
